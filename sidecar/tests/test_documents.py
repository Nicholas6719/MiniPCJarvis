"""Documents he can hand in: what gets written, where, and what comes back.

The dangerous failures here are not "it did not work". They are: a file
written outside his folders, a draft silently overwritten by a second one
with the same name, and a document that reports success while being empty.
All three are asserted.

Run: python tests/test_documents.py
"""
import asyncio
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("JARVIS_DB", os.path.join(tempfile.mkdtemp(), "docs.db"))

fails = []


def check(name, cond, detail=""):
    print(f"  {'PASS' if cond else 'FAIL'}  {name}  {detail if not cond else ''}")
    if not cond:
        fails.append(name)


ESSAY = """# The Water Cycle

Water moves in a closed loop. It is **never** created and never destroyed,
only moved.

## The stages

- Evaporation, driven by the sun
- Condensation into cloud
- Precipitation as rain or snow

1. First it rises
2. Then it gathers
3. Then it falls

| Stage | Where |
| --- | --- |
| Evaporation | Oceans |
| Precipitation | Everywhere |
"""


async def main() -> int:
    root = Path(tempfile.mkdtemp())
    for name in ("Documents", "Desktop", "Downloads", "Pictures"):
        (root / name).mkdir()
    # Patch the sandbox rather than the config: `config.set` SAVES, and a test
    # has no business rewriting where his real documents live.
    from tools import file_tools
    file_tools.roots = lambda: {"documents": root / "Documents",
                                "desktop": root / "Desktop",
                                "downloads": root / "Downloads",
                                "pictures": root / "Pictures"}
    from tools import documents as D

    print("\n-- a Word document from what the model writes --")
    r = await D.write_document("The Water Cycle", ESSAY, title="The Water Cycle")
    check("it is written", not r.get("error") and Path(r["path"]).exists(), r)
    check("...into his Documents folder",
          Path(r["path"]).parent == root / "Documents", r.get("path"))
    check("...with the headings, bullets and table the markdown asked for",
          r.get("headings") == 2 and r.get("bullets") == 6 and r.get("tables") == 1, r)
    check("...and it says how long it is", r.get("words", 0) > 20 and "words" in r.get("spoken", ""), r)

    from docx import Document
    doc = Document(r["path"])
    styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(p.text for p in doc.paragraphs)
    check("the heading is a real Word heading", any(s.startswith("Heading") or s == "Title"
                                                    for s in styles), styles[:6])
    check("the bullets are real list items", styles.count("List Bullet") == 3, styles)
    check("...and the numbered ones too", styles.count("List Number") == 3, styles)
    check("bold survives as bold, not as asterisks",
          "**" not in text and any(run.bold for p in doc.paragraphs for run in p.runs), text[:80])
    check("the table is a real table with its header",
          len(doc.tables) == 1 and doc.tables[0].rows[0].cells[0].text == "Stage",
          [t.rows[0].cells[0].text for t in doc.tables])
    check("wrapped lines are one paragraph, not two",
          "only moved." in text and "never destroyed,\nonly" not in text)

    print("\n-- a draft is never lost to a second one of the same name --")
    r2 = await D.write_document("The Water Cycle", "Another go entirely.")
    check("the same name makes a numbered sibling",
          r2["name"] == "The Water Cycle (2).docx" and Path(r["path"]).exists(), r2.get("name"))

    print("\n-- it cannot write outside his folders --")
    for bad in ("../../Windows/System32/evil", "C:/Windows/System32/evil",
                "..\\..\\..\\evil"):
        got = await D.write_document(bad, "no")
        wrote = got.get("path", "")
        check(f"{bad!r} stays inside the sandbox",
              not wrote or str(root) in wrote, wrote)
    check("an empty document is refused, not written",
          (await D.write_document("Empty", "   ")).get("error"), "")

    print("\n-- a spreadsheet --")
    s = await D.write_spreadsheet("Budget", rows=[["Item", "Cost"], ["Books", "120"],
                                                  ["Bus pass", "45.50"],
                                                  ["Total", "=SUM(B2:B3)"]])
    check("it is written", not s.get("error") and Path(s["path"]).exists(), s)
    check("...with the rows it was given", s.get("rows") == 4 and s.get("columns") == 2, s)
    from openpyxl import load_workbook
    ws = load_workbook(s["path"]).active
    check("a number written as text is stored as a number",
          isinstance(ws["B2"].value, int) and ws["B2"].value == 120, ws["B2"].value)
    check("...a decimal too", abs((ws["B3"].value or 0) - 45.5) < 1e-9, ws["B3"].value)
    check("a formula is kept as a formula", str(ws["B4"].value).startswith("=SUM"), ws["B4"].value)
    check("the header row is bold and frozen",
          ws["A1"].font.bold and ws.freeze_panes == "A2", (ws["A1"].font.bold, ws.freeze_panes))
    s2 = await D.write_spreadsheet("Marks", content="Subject,Grade\nMaths,A\nEnglish,B")
    check("CSV text works as well as rows", s2.get("rows") == 3, s2)
    s3 = await D.write_spreadsheet("Table", content="| A | B |\n| --- | --- |\n| 1 | 2 |")
    check("...and a markdown table", s3.get("rows") == 2, s3)
    check("an empty spreadsheet is refused",
          (await D.write_spreadsheet("Nothing")).get("error"), "")

    print("\n-- reading them back --")
    got = await D.read_document(r["path"])
    check("the Word document reads back", "Evaporation" in got.get("text", ""), got.get("error"))
    check("...with its table", "Oceans" in got.get("text", ""), "")
    got = await D.read_document(s["path"])
    check("the spreadsheet reads back", "Bus pass" in got.get("text", ""), got.get("error"))
    check("...naming its sheet", "[Sheet1]" in got.get("text", ""), got.get("text", "")[:40])
    note = await D.write_note("Reminder", "Chemistry test on Friday")
    check("a plain note is written", Path(note["path"]).exists() and note["path"].endswith(".txt"), note)
    got = await D.read_document("Reminder")
    check("a document can be found by name alone, without a path",
          "Chemistry" in got.get("text", ""), got.get("error"))
    check("a missing one says so", (await D.read_document("nothing at all")).get("error"), "")

    print("\n-- adding to one he already has --")
    a = await D.append_to_document(r["path"], "\n## Afterword\n\nAdded later.")
    check("it is added", not a.get("error"), a)
    got = await D.read_document(r["path"])
    check("...and the original is still there",
          "Evaporation" in got["text"] and "Added later." in got["text"], "")
    check("an unsupported format is refused, not mangled",
          (await D.append_to_document(s["path"], "x")).get("error"), "")

    print("\n-- the tools are registered --")
    D.register_all()
    from tools.registry import registry, Risk
    for name, risk in (("write_document", Risk.LOW), ("write_spreadsheet", Risk.LOW),
                       ("write_note", Risk.LOW), ("append_to_document", Risk.LOW),
                       ("read_document", Risk.SAFE)):
        t = registry.get(name)
        check(f"{name} is registered at {risk.name}", t is not None and t.risk is risk,
              t.risk if t else None)

    print(f"\n{'ALL PASS' if not fails else f'{len(fails)} FAILURES'}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
